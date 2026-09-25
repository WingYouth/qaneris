"""EXCEL-01A acceptance: ``Excel -> SQLite -> ScanSnapshot -> Neo4j`` on a real graph backend.

Run it with a real Neo4j reachable (``QANERIS_GRAPH_STORE=neo4j`` plus the ``QANERIS_NEO4J_*``
credentials) from the project root::

    python3 -m qaneris.scripts.acceptance.import_excel

What it proves, on a real endpoint rather than a double:

1. the deterministic ``orders`` workbook becomes a managed artifact
   (``source.xlsx`` / ``data.sqlite3`` / ``manifest.json``);
2. the materialized SQLite carries the expected tables, declared column types, row count, ISO dates
   and the preserved NULL;
3. the existing secure-datasource + ``DatabaseInitializer`` chain reaches datasource READY and an
   active READY ScanSnapshot, and Neo4j carries exactly one Database node for the datasource, the
   ``orders`` DataObject with five Fields, zero relationships and no duplicated node ids;
4. READY means verified: ``status == "READY"`` and ``neo4j_publication_verified is True`` together;
5. re-importing the same bytes under the same policy is idempotent (same import, datasource and
   snapshot, no second Database node, artifact not rewritten);
6. the same bytes under a different policy get a different artifact identity, and the finalized
   artifact of the first policy stays byte-for-byte, inode-for-inode untouched;
7. without a verifiable graph reader the import is NOT READY and leaves a FAILED manifest;
8. a publication that fails inside initialization, and a publication that is silently dropped, both
   leave a FAILED manifest with ``EXCEL_SCAN_FAILED`` - the dropped case is invisible to the scan and
   the snapshot, so only the publication check can catch it;
9. a workbook containing a formula is refused with ``EXCEL_FORMULA_UNSUPPORTED`` and never creates a
   datasource.

The run leaves its evidence under ``artifact_directory("acceptance")/excel-<run_id>``; the managed
artifacts themselves live under ``artifact_directory("ingestion")/excel``.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import sqlite3
import stat
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from qaneris.application.service import QanerisService
from qaneris.catalog import Catalog
from qaneris.common.artifacts import artifact_directory, artifact_root
from qaneris.graph import (
    Neo4jConfig,
    Neo4jGraphReader,
    NullGraphReader,
    NullGraphStore,
)
from qaneris.ingestion.excel import (
    DEFAULT_EXCEL_POLICY,
    ExcelErrorCode,
    ExcelImportRequest,
    ExcelIngestionError,
    ExcelIngestionPolicy,
    ExcelIngestionService,
    artifact_key,
)

ORDERS: list[list[object]] = [
    ["order_id", "order_date", "region", "amount", "status"],
    [1, dt.date(2026, 1, 5), "华东", 129.9, "completed"],
    [2, dt.date(2026, 1, 6), "华南", 899.5, "completed"],
    [3, dt.date(2026, 1, 7), "华北", 1899.25, "cancelled"],
    [4, dt.date(2026, 1, 8), "华中", 299.0, "completed"],
    [5, dt.date(2026, 2, 2), "西南", 6999.99, "completed"],
    [6, dt.date(2026, 2, 3), "华东", 3999.5, "pending"],
    [7, dt.date(2026, 2, 4), "华南", 2199.75, "completed"],
    [8, dt.date(2026, 2, 5), "华北", None, "pending"],
    [9, dt.date(2026, 3, 1), "华中", 599.5, "completed"],
    [10, dt.date(2026, 3, 2), "西南", 2499.9, "completed"],
]

EXPECTED_TYPES = {
    "order_id": "INTEGER",
    "order_date": "DATE",
    "region": "TEXT",
    "amount": "REAL",
    "status": "TEXT",
}

BUSINESS_TOKENS = ("华东", "华南", "华北", "华中", "西南", "completed", "cancelled", "pending")

# Differs from the default policy only by version: the case the artifact layout has to separate.
ALTERNATE_POLICY_VERSION = "excel-ingestion-acceptance-v2"

WORKSPACE = "default"


class FailingGraphStore:
    """A store whose publication raises, so initialization fails after the scan succeeded."""

    def ensure_schema(self) -> None:
        return None

    def replace_datasource_graph(self, graph: Any) -> None:
        raise RuntimeError("acceptance: simulated graph publication failure")


class DroppingGraphStore:
    """A store that accepts the publication and silently writes nothing."""

    def ensure_schema(self) -> None:
        return None

    def replace_datasource_graph(self, graph: Any) -> None:
        return None


def build_orders_workbook(path: Path, rows: list[list[object]] = ORDERS) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "orders"
    for row in rows:
        worksheet.append(row)
    for row_index in range(2, len(rows) + 1):
        worksheet.cell(row=row_index, column=2).number_format = "yyyy-mm-dd"
    workbook.save(path)
    return path


def probe_rows(marker: str) -> list[list[object]]:
    """The same shape as ``ORDERS`` with different bytes, so each probe owns its own artifact."""
    return [*ORDERS, [91, dt.date(2026, 3, 3), "华东", 12.5, marker]]


def build_formula_workbook(path: Path) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "orders"
    worksheet.append(["order_id", "amount"])
    worksheet.append([1, 100.0])
    worksheet.append([2, "=SUM(B2:B2)"])
    workbook.save(path)
    return path


def expected_directory(workbook: Path, policy: ExcelIngestionPolicy) -> Path:
    """The artifact directory the layout formula derives for these exact bytes and this policy."""
    workspace_digest = hashlib.sha256(WORKSPACE.encode("utf-8")).hexdigest()[:12]
    file_sha256 = hashlib.sha256(workbook.read_bytes()).hexdigest()
    return (
        artifact_directory("ingestion") / "excel" / workspace_digest / artifact_key(policy) / file_sha256
    )


def neo4j_reader() -> Neo4jGraphReader:
    config = Neo4jConfig.from_environment()
    return Neo4jGraphReader(config.uri, config.username, config.password, database=config.database)


def count_database_nodes(datasource_id: str) -> int:
    config = Neo4jConfig.from_environment()
    from neo4j import GraphDatabase

    driver = GraphDatabase.driver(config.uri, auth=(config.username, config.password))
    try:
        with driver.session(database=config.database) as session:
            record = session.run(
                "MATCH (d:Database {datasource_id: $datasource_id}) RETURN count(d) AS total",
                datasource_id=datasource_id,
            ).single()
    finally:
        driver.close()
    return int(record["total"])


def read_graph(datasource_id: str) -> dict[str, object]:
    config = Neo4jConfig.from_environment()
    from neo4j import GraphDatabase

    driver = GraphDatabase.driver(config.uri, auth=(config.username, config.password))
    try:
        with driver.session(database=config.database) as session:
            row = session.run(
                "MATCH (d:Database {datasource_id: $datasource_id}) "
                "OPTIONAL MATCH (d)-[*1..]->(n) "
                "WITH d, collect(DISTINCT n) AS nodes "
                "OPTIONAL MATCH (a:DataObject)-[r:RELATES_TO]->(b:DataObject) "
                "WHERE a.datasource_id=$datasource_id AND b.datasource_id=$datasource_id "
                "RETURN count(DISTINCT d) AS databases, "
                "d.scan_version AS scan_version, "
                "[node IN nodes WHERE node:DataObject | node.name] AS object_names, "
                "size([node IN nodes WHERE node:Field]) AS fields, "
                "count(DISTINCT r) AS relationships",
                datasource_id=datasource_id,
            ).single()
            integrity = session.run(
                "MATCH (d:Database {datasource_id: $datasource_id})-[*0..]->(n) "
                "WITH collect(DISTINCT n) AS nodes "
                "UNWIND nodes AS n WITH nodes, n.id AS id, count(*) AS occurrences "
                "WITH nodes, sum(CASE WHEN occurrences > 1 THEN 1 ELSE 0 END) AS duplicates "
                "RETURN duplicates, size([n IN nodes WHERE n.id IS NULL]) AS missing_ids",
                datasource_id=datasource_id,
            ).single()
    finally:
        driver.close()
    return {
        "databases": int(row["databases"]),
        "scan_version": row["scan_version"],
        "object_names": sorted(name for name in row["object_names"] if name),
        "fields": int(row["fields"]),
        "relationships": int(row["relationships"]),
        "duplicates": int(integrity["duplicates"]),
        "missing_ids": int(integrity["missing_ids"]),
    }


def read_sqlite(path: Path) -> dict[str, object]:
    connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        tables = [
            name
            for (name,) in connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
            ).fetchall()
        ]
        declared = {
            name: declared_type.upper()
            for _, name, declared_type, *_ in connection.execute("PRAGMA table_info(orders)")
        }
        rows = connection.execute('SELECT count(*) FROM "orders"').fetchone()[0]
        nulls = connection.execute(
            'SELECT count(*) FROM "orders" WHERE "amount" IS NULL'
        ).fetchone()[0]
        sample = connection.execute(
            'SELECT "order_date", "amount", "status" FROM "orders" WHERE "order_id" = 5'
        ).fetchone()
        iso_text = connection.execute(
            'SELECT "order_date" FROM "orders" WHERE "order_id" = 5'
        ).fetchone()[0]
    finally:
        connection.close()
    return {
        "tables": tables,
        "declared_types": declared,
        "rows": int(rows),
        "null_amounts": int(nulls),
        "sample": sample,
        "iso_text": iso_text,
    }


def artifact_modes(directory: Path) -> dict[str, str]:
    return {
        str(path.relative_to(directory)): oct(stat.S_IMODE(path.stat().st_mode))
        for path in sorted(directory.rglob("*"))
    }


def fingerprint(directory: Path) -> dict[str, tuple[int, int, int, str]]:
    """Identity of the finalized artifacts: inode, mtime, size and content hash."""
    return {
        name: (
            (directory / name).stat().st_ino,
            (directory / name).stat().st_mtime_ns,
            (directory / name).stat().st_size,
            hashlib.sha256((directory / name).read_bytes()).hexdigest(),
        )
        for name in ("source.xlsx", "data.sqlite3", "manifest.json")
        if (directory / name).is_file()
    }


def manifest_at(directory: Path) -> dict[str, Any] | None:
    path = directory / "manifest.json"
    if not path.is_file():
        return None
    return json.loads(path.read_text("utf-8"))


def managed_artifact_directories() -> list[Path]:
    """Every managed artifact directory this workspace currently has, for the run summary."""
    workspace_digest = hashlib.sha256(WORKSPACE.encode("utf-8")).hexdigest()[:12]
    root = artifact_directory("ingestion") / "excel" / workspace_digest
    if not root.is_dir():
        return []
    return sorted(path.parent for path in root.glob("*/**/manifest.json"))


def attempt(service: QanerisService, request: ExcelImportRequest) -> ExcelIngestionError | None:
    """Import and return the refusal, or ``None`` when the import claimed READY."""
    try:
        service.import_excel(request)
    except ExcelIngestionError as error:
        return error
    return None


def main() -> int:
    """One linear acceptance narrative, kept in a single function so the story is readable."""
    run_id = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    run_root = artifact_directory("acceptance") / f"excel-{run_id}"
    run_root.mkdir(parents=True, exist_ok=False)

    workbook_path = build_orders_workbook(run_root / "orders.xlsx")
    probe_paths = {
        marker: build_orders_workbook(run_root / f"{marker}.xlsx", probe_rows(marker))
        for marker in ("no-graph", "scan-failure", "dropped-publication")
    }
    formula_path = build_formula_workbook(run_root / "formula.xlsx")

    # ---- 1. the real chain: import, verify, replay -------------------------------------
    service = QanerisService(Catalog(run_root / "catalog.db"))
    request = ExcelImportRequest(
        file_path=str(workbook_path), datasource_name=f"excel-acceptance-{run_id}"
    )

    result = service.import_excel(request)

    directory = Path(result.artifact_directory)
    sqlite_state = read_sqlite(Path(result.sqlite_path))
    graph = read_graph(result.datasource_id)
    manifest = manifest_at(directory) or {}
    datasource, _ = service.catalog.get_datasource(result.datasource_id)
    snapshot = service.catalog.get_active_snapshot(result.datasource_id)
    modes = artifact_modes(directory)
    manifest_text = json.dumps(manifest, ensure_ascii=False)

    directory_before = fingerprint(directory)
    replay = service.import_excel(request)
    graph_after_replay = read_graph(result.datasource_id)
    directory_after = fingerprint(directory)

    # ---- 2. same bytes, different policy: new identity, old artifact untouched ---------
    alternate_policy = replace(DEFAULT_EXCEL_POLICY, version=ALTERNATE_POLICY_VERSION)
    # The service-level entry point is used here because a non-default policy has no interface
    # layer yet (the CLI exposes the default policy only; Web Upload is EXCEL-01C): the subject of
    # this section is the artifact identity.
    alternate = ExcelIngestionService(
        QanerisService(Catalog(run_root / "catalog-alternate.db")), policy=alternate_policy
    )
    alternate_result = alternate.import_excel(
        ExcelImportRequest(
            file_path=str(workbook_path), datasource_name=f"excel-acceptance-v2-{run_id}"
        )
    )
    directory_after_policy_change = fingerprint(directory)

    # ---- 3. no verifiable graph reader -------------------------------------------------
    no_graph_service = QanerisService(
        Catalog(run_root / "catalog-no-graph.db"),
        graph_store=NullGraphStore(),
        graph_reader=NullGraphReader(),
    )
    no_graph_error = attempt(
        no_graph_service,
        ExcelImportRequest(
            file_path=str(probe_paths["no-graph"]), datasource_name=f"excel-no-graph-{run_id}"
        ),
    )

    # ---- 4. publication raises inside initialization -----------------------------------
    scan_failure_service = QanerisService(
        Catalog(run_root / "catalog-scan-failure.db"),
        graph_store=FailingGraphStore(),
        graph_reader=neo4j_reader(),
    )
    scan_failure_error = attempt(
        scan_failure_service,
        ExcelImportRequest(
            file_path=str(probe_paths["scan-failure"]),
            datasource_name=f"excel-scan-failure-{run_id}",
        ),
    )

    # ---- 5. publication silently dropped ------------------------------------------------
    dropped_service = QanerisService(
        Catalog(run_root / "catalog-dropped.db"),
        graph_store=DroppingGraphStore(),
        graph_reader=neo4j_reader(),
    )
    dropped_error = attempt(
        dropped_service,
        ExcelImportRequest(
            file_path=str(probe_paths["dropped-publication"]),
            datasource_name=f"excel-dropped-{run_id}",
        ),
    )

    # ---- 6. a rejected workbook creates no datasource ----------------------------------
    formula_error = attempt(
        service,
        ExcelImportRequest(file_path=str(formula_path), datasource_name=f"excel-formula-{run_id}"),
    )
    datasources_after_formula = service.list_datasources(WORKSPACE)

    no_graph_manifest = manifest_at(expected_directory(probe_paths["no-graph"], DEFAULT_EXCEL_POLICY))
    scan_failure_manifest = manifest_at(
        expected_directory(probe_paths["scan-failure"], DEFAULT_EXCEL_POLICY)
    )
    dropped_manifest = manifest_at(
        expected_directory(probe_paths["dropped-publication"], DEFAULT_EXCEL_POLICY)
    )
    probe_directories = [
        expected_directory(path, DEFAULT_EXCEL_POLICY) for path in probe_paths.values()
    ]
    managed_directories = managed_artifact_directories()

    checks = {
        # 1. real chain
        "Import READY": result.status == "READY",
        "Publication verified": result.neo4j_publication_verified is True,
        "READY implies verified": not (
            result.status == "READY" and not result.neo4j_publication_verified
        ),
        "Artifact layout": sorted(modes) == ["data.sqlite3", "manifest.json", "source.xlsx"],
        "Artifact privacy": all(mode == "0o600" for mode in modes.values())
        and oct(stat.S_IMODE(directory.stat().st_mode)) == "0o700",
        "Artifact levels private": all(
            oct(stat.S_IMODE(level.stat().st_mode)) == "0o700"
            for level in (
                directory,
                directory.parent,
                directory.parent.parent,
                directory.parent.parent.parent,
            )
        ),
        "Artifact directory is content-addressed": directory
        == expected_directory(workbook_path, DEFAULT_EXCEL_POLICY),
        "Artifact path has no business filename": all(
            ".xlsx" not in part and "订单" not in part
            for part in directory.relative_to(artifact_root()).parts
        ),
        "SQLite tables": sqlite_state["tables"] == ["orders"],
        "SQLite declared types": sqlite_state["declared_types"] == EXPECTED_TYPES,
        "SQLite row count": sqlite_state["rows"] == len(ORDERS) - 1,
        "SQLite NULL preserved": sqlite_state["null_amounts"] == 1,
        "SQLite DATE stored as ISO text": sqlite_state["iso_text"] == "2026-02-02",
        "SQLite REAL not truncated": isinstance(sqlite_state["sample"][1], float)
        and abs(sqlite_state["sample"][1] - 6999.99) < 1e-9,
        "Manifest READY": manifest.get("status") == "READY",
        "Manifest redaction": not any(token in manifest_text for token in BUSINESS_TOKENS),
        "Manifest shape": manifest.get("sheets", [{}])[0].get("row_count") == len(ORDERS) - 1,
        "Datasource READY": datasource.status == "ready",
        "Active snapshot matches": snapshot is not None
        and snapshot.id == result.snapshot_id
        and snapshot.version == result.scan_version,
        "Graph: one Database node": graph["databases"] == 1,
        "Graph: orders DataObject": graph["object_names"] == ["orders"],
        "Graph: five Fields": graph["fields"] == len(EXPECTED_TYPES),
        "Graph: no relationships": graph["relationships"] == 0,
        "Graph: scan_version published": graph["scan_version"] == result.scan_version,
        "Graph integrity": graph["duplicates"] == 0 and graph["missing_ids"] == 0,
        # 5. same file + same policy
        "Idempotent import id": replay.import_id == result.import_id,
        "Idempotent datasource": replay.datasource_id == result.datasource_id,
        "Idempotent snapshot": replay.snapshot_id == result.snapshot_id,
        "Idempotent graph": graph_after_replay["databases"] == 1
        and graph_after_replay["fields"] == graph["fields"],
        "Idempotent artifact": directory_after == directory_before,
        # 6. same file + different policy
        "Policy: new artifact directory": alternate_result.artifact_directory
        != result.artifact_directory,
        "Policy: new import id": alternate_result.import_id != result.import_id,
        "Policy: alternate READY + verified": alternate_result.status == "READY"
        and alternate_result.neo4j_publication_verified is True,
        "Policy: old artifact untouched": directory_after_policy_change == directory_before,
        "Policy: old manifest still READY": manifest.get("status") == "READY",
        "Policy: old import still idempotent": service.import_excel(request).import_id
        == result.import_id,
        # 7. no verifiable graph reader
        "No graph: NOT READY": no_graph_error is not None,
        "No graph: EXCEL_SCAN_FAILED": no_graph_error is not None
        and no_graph_error.error_code == ExcelErrorCode.SCAN_FAILED.value,
        "No graph: FAILED manifest": no_graph_manifest is not None
        and no_graph_manifest["status"] == "FAILED"
        and no_graph_manifest["error_code"] == ExcelErrorCode.SCAN_FAILED.value,
        # 8a. publication failure during initialization
        "Scan failure: EXCEL_SCAN_FAILED": scan_failure_error is not None
        and scan_failure_error.error_code == ExcelErrorCode.SCAN_FAILED.value,
        "Scan failure: FAILED manifest": scan_failure_manifest is not None
        and scan_failure_manifest["status"] == "FAILED"
        and scan_failure_manifest["error_code"] == ExcelErrorCode.SCAN_FAILED.value,
        "Scan failure: no fabricated snapshot": scan_failure_manifest is not None
        and scan_failure_manifest["snapshot_id"] is None,
        # 8b. publication silently dropped
        "Dropped publication: EXCEL_SCAN_FAILED": dropped_error is not None
        and dropped_error.error_code == ExcelErrorCode.SCAN_FAILED.value,
        "Dropped publication: FAILED manifest": dropped_manifest is not None
        and dropped_manifest["status"] == "FAILED"
        and dropped_manifest["error_code"] == ExcelErrorCode.SCAN_FAILED.value,
        "Dropped publication: nothing in Neo4j": dropped_manifest is not None
        and dropped_manifest["datasource_id"] is not None
        and count_database_nodes(str(dropped_manifest["datasource_id"])) == 0,
        "Dropped publication: snapshot still committed": dropped_manifest is not None
        and dropped_manifest["snapshot_id"] is not None,
        # every failed attempt kept its finalized source
        "Failures keep the source": all(
            (probe / "source.xlsx").is_file()
            for probe in (probe_directories + [expected_directory(formula_path, DEFAULT_EXCEL_POLICY)])
        ),
        # 9. formula workbook
        "Formula rejected": formula_error is not None
        and formula_error.error_code == ExcelErrorCode.FORMULA_UNSUPPORTED.value,
        "Formula created no datasource": len(datasources_after_formula) == 1,
    }

    for name, passed in checks.items():
        print(f"[{'PASS' if passed else 'FAIL'}] {name}")
    print(f"Import id      : {result.import_id}")
    print(f"Datasource id  : {result.datasource_id}")
    print(f"Snapshot id    : {result.snapshot_id} (scan_version={result.scan_version})")
    print(f"Artifact       : {directory}")
    print(f"Artifact (v2)  : {alternate_result.artifact_directory}")
    print(f"Managed dirs   : {len(managed_directories)}")
    print(f"Evidence       : {run_root}")
    failed = [name for name, passed in checks.items() if not passed]
    print(
        f"Result         : {'PASS' if not failed else 'FAIL'} "
        f"({len(checks) - len(failed)}/{len(checks)})"
    )
    return 0 if not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
