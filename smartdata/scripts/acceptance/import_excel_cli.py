"""EXCEL-01B acceptance: the real ``smartdata source import-excel`` CLI on a real Neo4j.

Run it from the project root with a real graph backend exported (``SMARTDATA_GRAPH_STORE=neo4j``
plus the ``SMARTDATA_NEO4J_*`` credentials, for example ``set -a && . ./.env.acceptance && set +a``)::

    python3 -m smartdata.scripts.acceptance.import_excel_cli

Unlike the EXCEL-01A acceptance, which drives ``SmartDataService`` in-process, this one runs the
installed console script as a subprocess, so what is verified is the product surface itself. It
proves:

1. ``smartdata source import-excel FILE --name NAME --workspace default`` exits 0, prints the
   identity block and leaves a datasource that Neo4j really carries at the reported revision;
2. ``--json`` exits 0 with a single parseable JSON document on stdout, and the document says
   ``READY`` with ``neo4j_publication_verified`` true - it never exposes internal file locations;
3. repeating the same command is idempotent: same ``import_id``, ``datasource_id`` and
   ``snapshot_id``, with the managed artifact untouched;
4. a workbook containing a formula exits 1 with ``EXCEL_FORMULA_UNSUPPORTED``, a safe structured
   error, no formula body and no business cell value, and creates no datasource;
5. a file that is not a readable workbook exits 1 with ``EXCEL_INVALID_CONTAINER``.

The run keeps an isolated catalog inside its evidence directory so repeated runs stay reproducible;
everything else - the Neo4j backend and the artifact root - comes from the environment unchanged.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import sqlite3
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from smartdata.common.artifacts import artifact_directory, project_root
from smartdata.contracts.profile import ScanStatus
from smartdata.graph import Neo4jConfig
from smartdata.scripts.acceptance.summary import write_summary

WORKSPACE = "default"
DATASOURCE_NAME = "Roadshow Excel Orders"

# A shape of its own so this acceptance owns its artifact identity instead of reusing the one the
# EXCEL-01A acceptance already materialized for its five-column ``orders`` workbook.
CLI_ORDERS: list[list[object]] = [
    ["order_id", "order_date", "region", "amount", "channel", "status"],
    [1, dt.date(2026, 1, 5), "华东", 129.9, "online", "completed"],
    [2, dt.date(2026, 1, 6), "华南", 899.5, "store", "completed"],
    [3, dt.date(2026, 1, 7), "华北", 1899.25, "online", "cancelled"],
    [4, dt.date(2026, 1, 8), "华中", 299.0, "store", "completed"],
    [5, dt.date(2026, 2, 2), "西南", 6999.99, "online", "completed"],
    [6, dt.date(2026, 2, 3), "华东", 3999.5, "store", "pending"],
    [7, dt.date(2026, 2, 4), "华南", 2199.75, "online", "completed"],
    [8, dt.date(2026, 2, 5), "华北", None, "store", "pending"],
    [9, dt.date(2026, 3, 1), "华中", 599.5, "online", "completed"],
    [10, dt.date(2026, 3, 2), "西南", 2499.9, "store", "completed"],
]

EXPECTED_ROW_COUNT = len(CLI_ORDERS) - 1
EXPECTED_COLUMN_COUNT = len(CLI_ORDERS[0])

# A formula body or a stored business value must never appear in CLI output.
FORBIDDEN_TOKENS = ("=SUM(", "SUM(B2", "华东", "华南", "completed", "cancelled", "online")

HUMAN_LINES = ("Datasource: ", "Snapshot: ", "Scan version: ", "Sheets: ", "Artifact: ")


def build_workbook(path: Path, rows: list[list[object]], date_column: int | None = 2) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "orders"
    for row in rows:
        worksheet.append(row)
    if date_column is not None:
        for row_index in range(2, len(rows) + 1):
            worksheet.cell(row=row_index, column=date_column).number_format = "yyyy-mm-dd"
    workbook.save(path)
    return path


def build_formula_workbook(path: Path) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "orders"
    worksheet.append(["order_id", "amount"])
    worksheet.append([1, 100.0])
    worksheet.append([2, "=SUM(B2:B2)"])
    workbook.save(path)
    return path


def cli_command() -> list[str]:
    """Prefer the installed console script: that is the product surface being accepted."""
    script = Path(sys.executable).parent / "smartdata"
    if script.is_file():
        return [str(script)]
    return [sys.executable, "-m", "smartdata.cli.main"]


def run_cli(
    arguments: list[str], *, cwd: Path, environment: dict[str, str]
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [*cli_command(), *arguments],
        cwd=cwd,
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )


def field(output: str, label: str) -> str | None:
    match = re.search(rf"^{re.escape(label)}: (.+)$", output, re.MULTILINE)
    return match.group(1).strip() if match else None


def read_catalog(path: Path) -> dict[str, Any]:
    connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        datasources = [
            {"id": row[0], "name": row[1], "status": row[2]}
            for row in connection.execute(
                "SELECT id, name, status FROM datasource ORDER BY id"
            )
        ]
        snapshots = [
            {"id": row[0], "datasource_id": row[1], "version": row[2], "status": row[3],
             "active": row[4]}
            for row in connection.execute(
                "SELECT id, datasource_id, version, status, active "
                "FROM scan_snapshot ORDER BY id"
            )
        ]
    finally:
        connection.close()
    return {"datasources": datasources, "snapshots": snapshots}


def read_graph(datasource_id: str) -> dict[str, Any]:
    """Read the published graph directly, so the evidence does not come from the same code path."""
    config = Neo4jConfig.from_environment()
    from neo4j import GraphDatabase

    driver = GraphDatabase.driver(config.uri, auth=(config.username, config.password))
    try:
        with driver.session(database=config.database) as session:
            row = session.run(
                "MATCH (d:Database {datasource_id: $datasource_id}) "
                "OPTIONAL MATCH (d)-[:CONTAINS]->(o:DataObject) "
                "OPTIONAL MATCH (o)-[:HAS_FIELD]->(f:Field) "
                "OPTIONAL MATCH (a:DataObject {datasource_id: $datasource_id})"
                "-[r:RELATES_TO]->(b:DataObject) "
                "RETURN count(DISTINCT d) AS databases, "
                "max(d.scan_version) AS scan_version, "
                "count(DISTINCT o) AS objects, "
                "count(DISTINCT f) AS fields, "
                "count(DISTINCT r) AS relationships, "
                "[name IN collect(DISTINCT o.name) WHERE name IS NOT NULL] AS object_names",
                datasource_id=datasource_id,
            ).single()
    finally:
        driver.close()
    return {
        "databases": int(row["databases"]),
        "scan_version": row["scan_version"],
        "objects": int(row["objects"]),
        "fields": int(row["fields"]),
        "relationships": int(row["relationships"]),
        "object_names": sorted(row["object_names"]),
    }


def fingerprint(directory: Path) -> dict[str, tuple[int, int, str]]:
    return {
        name: (
            (directory / name).stat().st_ino,
            (directory / name).stat().st_mtime_ns,
            hashlib.sha256((directory / name).read_bytes()).hexdigest(),
        )
        for name in ("source.xlsx", "data.sqlite3", "manifest.json")
        if (directory / name).is_file()
    }


def main() -> int:
    run_id = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    run_root = artifact_directory("acceptance") / f"excel-cli-{run_id}"
    run_root.mkdir(parents=True, exist_ok=False)

    try:
        Neo4jConfig.from_environment()
    except ValueError as error:
        print(f"[FAIL] Neo4j is not configured: {error}")
        print("       export SMARTDATA_GRAPH_STORE=neo4j and the SMARTDATA_NEO4J_* credentials")
        return 1
    if os.getenv("SMARTDATA_GRAPH_STORE", "neo4j").strip().lower() == "null":
        print("[FAIL] SMARTDATA_GRAPH_STORE=null: READY requires a verified publication")
        return 1

    workbook = build_workbook(run_root / "orders.xlsx", CLI_ORDERS)
    formula = build_formula_workbook(run_root / "formula.xlsx")
    missing = run_root / "does-not-exist.xlsx"

    catalog_path = run_root / "smartdata.db"
    environment = {**os.environ, "SMARTDATA_CATALOG": str(catalog_path)}
    cwd = project_root()

    # ---- 1. human run ------------------------------------------------------------------
    human = run_cli(
        ["source", "import-excel", str(workbook), "--name", DATASOURCE_NAME,
         "--workspace", WORKSPACE],
        cwd=cwd,
        environment=environment,
    )
    artifact_directory_path = field(human.stdout, "Artifact")
    directory = Path(artifact_directory_path) if artifact_directory_path else None

    # ---- 2. JSON runs, the second of which proves idempotency ---------------------------
    first_json = run_cli(
        ["source", "import-excel", str(workbook), "--name", DATASOURCE_NAME,
         "--workspace", WORKSPACE, "--json"],
        cwd=cwd,
        environment=environment,
    )
    payload = _load(first_json.stdout)
    before_replay = fingerprint(directory) if directory else {}
    replay = run_cli(
        ["source", "import-excel", str(workbook), "--name", DATASOURCE_NAME,
         "--workspace", WORKSPACE, "--json"],
        cwd=cwd,
        environment=environment,
    )
    replay_payload = _load(replay.stdout)
    after_replay = fingerprint(directory) if directory else {}

    datasource_id = str(payload.get("datasource_id", ""))
    graph = read_graph(datasource_id) if datasource_id else {}
    catalog = read_catalog(catalog_path)

    # ---- 3. negative: a formula workbook ------------------------------------------------
    datasources_before = len(catalog["datasources"])
    negative = run_cli(
        ["source", "import-excel", str(formula), "--json"], cwd=cwd, environment=environment
    )
    negative_payload = _load(negative.stdout)
    datasources_after = len(read_catalog(catalog_path)["datasources"])

    # ---- 4. negative: not a readable workbook ------------------------------------------
    missing_run = run_cli(
        ["source", "import-excel", str(missing), "--json"], cwd=cwd, environment=environment
    )
    missing_payload = _load(missing_run.stdout)

    negative_streams = (
        f"{negative.stdout} {negative.stderr} {missing_run.stdout} "
        f"{missing_run.stderr} {first_json.stdout} {replay.stdout}"
    )

    checks = {
        # 1. human product surface
        "Human exit 0": human.returncode == 0,
        "Human import ready": "[PASS] Excel import ready" in human.stdout,
        "Human name echoed": field(human.stdout, "Name") == DATASOURCE_NAME,
        "Human identity block": all(line in human.stdout for line in HUMAN_LINES),
        "Human sheets": field(human.stdout, "Sheets") == "1",
        "Human artifact is the managed directory": directory is not None
        and directory.name
        == hashlib.sha256(workbook.read_bytes()).hexdigest()
        and directory.is_dir(),
        # 2. machine surface
        "JSON exit 0": first_json.returncode == 0,
        "JSON stdout is one document": payload != {},
        "JSON status READY": payload.get("status") == "READY",
        "JSON publication verified": payload.get("neo4j_publication_verified") is True,
        "JSON ids present": bool(payload.get("datasource_id"))
        and bool(payload.get("snapshot_id")),
        "JSON scan version 1": payload.get("scan_version") == 1,
        "JSON original filename": payload.get("original_filename") == "orders.xlsx",
        "JSON sheet shape": payload.get("sheets", [{}])[0].get("row_count")
        == EXPECTED_ROW_COUNT
        and payload.get("sheets", [{}])[0].get("column_count") == EXPECTED_COLUMN_COUNT,
        "JSON hides internal locations": "data.sqlite3" not in first_json.stdout
        and "manifest.json" not in first_json.stdout,
        "JSON artifact directory agrees": directory is not None
        and payload.get("artifact_directory") == str(directory),
        # 3. idempotency through the CLI
        "Replay exit 0": replay.returncode == 0,
        "Replay same import": replay_payload.get("import_id") == payload.get("import_id"),
        "Replay same datasource": replay_payload.get("datasource_id") == datasource_id,
        "Replay same snapshot": replay_payload.get("snapshot_id") == payload.get("snapshot_id"),
        "Replay same scan version": replay_payload.get("scan_version")
        == payload.get("scan_version"),
        "Replay leaves the artifact untouched": before_replay == after_replay,
        "Replay adds no datasource": len(catalog["datasources"]) == 1,
        # 4. the published graph really carries this datasource at that revision
        "Graph has one Database node": graph.get("databases") == 1,
        "Graph scan version matches": graph.get("scan_version") == payload.get("scan_version"),
        "Graph has the DataObject": graph.get("object_names") == ["orders"],
        "Graph field count": graph.get("fields") == EXPECTED_COLUMN_COUNT,
        "Graph has no invented relationships": graph.get("relationships") == 0,
        # 5. the catalog agrees with the reported identity
        "Catalog datasource ready": catalog["datasources"]
        == [{"id": datasource_id, "name": DATASOURCE_NAME, "status": "ready"}],
        "Catalog snapshot ready": len(catalog["snapshots"]) == 1
        and catalog["snapshots"][0]["id"] == payload.get("snapshot_id")
        and catalog["snapshots"][0]["datasource_id"] == datasource_id
        and catalog["snapshots"][0]["status"] == ScanStatus.READY.value
        and catalog["snapshots"][0]["version"] == payload.get("scan_version")
        and catalog["snapshots"][0]["active"] == 1,
        # 6. negative acceptance
        "Formula exit 1": negative.returncode == 1,
        "Formula code preserved": negative_payload.get("error", {}).get("code")
        == "EXCEL_FORMULA_UNSUPPORTED",
        "Formula structured error": negative_payload.get("status") == "operation_failed"
        and negative_payload.get("error", {}).get("type") == "ExcelIngestionError"
        and negative_payload.get("error", {}).get("sheet") == "orders"
        and negative_payload.get("error", {}).get("coordinate") == "B3",
        "Formula creates no datasource": datasources_after == datasources_before,
        "Missing file exit 1": missing_run.returncode == 1,
        "Missing file code preserved": missing_payload.get("error", {}).get("code")
        == "EXCEL_INVALID_CONTAINER",
        "No formula body or business value leaked": not any(
            token in negative_streams for token in FORBIDDEN_TOKENS
        ),
        "No secret leaked": not _leaked_secret(negative_streams),
    }

    for name, passed in checks.items():
        print(f"[{'PASS' if passed else 'FAIL'}] {name}")

    write_summary(
        checks,
        stage="cli_import",
        datasource_id=datasource_id,
        snapshot_id=payload.get("snapshot_id"),
        scan_version=payload.get("scan_version"),
        run_root=str(run_root),
    )

    print(f"CLI            : {' '.join(cli_command())}")
    print(f"Import id      : {payload.get('import_id')}")
    print(f"Datasource id  : {datasource_id}")
    print(f"Snapshot id    : {payload.get('snapshot_id')} "
          f"(scan_version={payload.get('scan_version')})")
    print(f"Neo4j verified : {payload.get('neo4j_publication_verified')}")
    print(f"Artifact       : {directory}")
    print(f"Evidence       : {run_root}")
    failed = [name for name, passed in checks.items() if not passed]
    print(
        f"Result         : {'PASS' if not failed else 'FAIL'} "
        f"({len(checks) - len(failed)}/{len(checks)})"
    )
    return 0 if not failed else 1


def _load(output: str) -> dict[str, Any]:
    """Parse the machine output; the whole stdout must be exactly one JSON document."""
    try:
        parsed = json.loads(output)
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _leaked_secret(text: str) -> bool:
    """True when any credential-looking environment value reached the output."""
    markers = ("PASSWORD", "SECRET", "TOKEN", "API_KEY", "CREDENTIAL", "PRIVATE_KEY")
    return any(
        value in text
        for name, value in os.environ.items()
        if value
        and len(value) >= 4
        and any(marker in name.upper() for marker in markers)
    )


if __name__ == "__main__":
    raise SystemExit(main())
