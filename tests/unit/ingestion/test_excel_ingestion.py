"""Excel ingestion unit tests (EXCEL-01A core ingestion).

These tests lock the deterministic contract from ``docs/excel-ingestion.md``: container preflight,
sheet inclusion, header/identifier rules, formula and merged-range rejection, resource limits, the
type lattice, SQLite materialization, artifact redaction, idempotency and the READY definition.

READY requires a verified Neo4j publication, so the tests never lower that definition to make an
import convenient: the READY paths inject a graph backend that really records what
``DatabaseInitializer`` published and really answers the read request, and the failure paths use a
backend that drops the write or raises.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import sqlite3
import zipfile
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from qaneris.application.service import QanerisService
from qaneris.catalog import Catalog
from qaneris.common.errors import GraphUnavailableError, QuerySafetyError
from qaneris.graph import (
    GraphDatasource,
    GraphStructure,
    GraphStructureRequest,
    NullGraphReader,
)
from qaneris.ingestion.excel import (
    DEFAULT_EXCEL_POLICY,
    ColumnPlan,
    ExcelErrorCode,
    ExcelImportRequest,
    ExcelIngestionError,
    ExcelIngestionPolicy,
    ExcelIngestionService,
    SQLiteMaterializer,
    TableWrite,
    artifact_key,
    classify_cell,
    coerce_cell,
    infer_column_type,
    iter_table_rows,
    policy_digest,
    quote_identifier,
    read_container,
    validate_workbook,
)
from qaneris.ingestion.excel.types import CellKind, ColumnType, UnsupportedCellValue

# --------------------------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------------------------


class FakeGraph:
    """In-memory stand-in for the published enterprise data graph.

    The store records exactly what ``DatabaseInitializer`` published and the reader answers the read
    request from that record, so publication verification is a real observation: a store that
    silently drops the write, or a reader that cannot answer, fails the import. Flip ``drop_writes``
    to publish nothing without raising, or ``unavailable`` to behave like an unreachable Neo4j.
    """

    def __init__(self) -> None:
        self.published: dict[tuple[str, str], int] = {}
        self.ensure_schema_calls = 0
        self.drop_writes = False
        self.unavailable = False

    # -- GraphStore -------------------------------------------------------------------------
    def ensure_schema(self) -> None:
        self.ensure_schema_calls += 1

    def replace_datasource_graph(self, graph: Any) -> None:
        if self.drop_writes:
            return
        self.published[(graph.workspace_id, graph.datasource_id)] = graph.scan_version

    # -- GraphReader ------------------------------------------------------------------------
    def read_structure(self, request: GraphStructureRequest) -> GraphStructure:
        if self.unavailable:
            raise GraphUnavailableError("simulated Neo4j outage")
        return GraphStructure(
            datasources=[
                GraphDatasource(
                    node_id=f"node::{datasource_id}",
                    datasource_id=datasource_id,
                    workspace_id=workspace,
                    name=datasource_id,
                    scan_version=version,
                )
                for (workspace, datasource_id), version in sorted(self.published.items())
                if workspace == request.workspace_id
                and (request.datasource_id is None or datasource_id == request.datasource_id)
            ]
        )


class FailingGraphStore:
    """Publishes nothing and raises: the publication step of initialization fails."""

    def ensure_schema(self) -> None:
        return None

    def replace_datasource_graph(self, graph: Any) -> None:
        raise RuntimeError("simulated graph publication failure")


def build_workbook(
    sheets: dict[str, list[list[Any]]],
    *,
    order: list[str] | None = None,
    number_formats: dict[str, dict[str, str]] | None = None,
    sheet_states: dict[str, str] | None = None,
    merges: dict[str, list[str]] | None = None,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title in order or list(sheets):
        worksheet = workbook.create_sheet(title)
        for row in sheets[title]:
            worksheet.append(row)
        for coordinate, code in (number_formats or {}).get(title, {}).items():
            worksheet[coordinate].number_format = code
        for reference in (merges or {}).get(title, []):
            worksheet.merge_cells(reference)
        if title in (sheet_states or {}):
            worksheet.sheet_state = (sheet_states or {})[title]
    return _save(workbook)


def _save(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_source(tmp_path: Path, data: bytes, name: str = "orders.xlsx") -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


def simple_rows() -> list[list[Any]]:
    return [["order_id", "region"], [1, "east"], [2, "west"]]


def naive_datetime(
    year: int, month: int, day: int, hour: int = 0, minute: int = 0, second: int = 0
) -> dt.datetime:
    """Excel stores naive local timestamps: the file carries no timezone, so the tests do not either."""
    return dt.datetime.combine(dt.date(year, month, day), dt.time(hour, minute, second))


def ingestion(tmp_path: Path, application: QanerisService, **changes) -> ExcelIngestionService:
    return ExcelIngestionService(
        application,
        policy=changes.pop("policy", None),
        artifact_directory_root=changes.pop("artifact_root", tmp_path / "artifacts"),
        **changes,
    )


def application(
    tmp_path: Path,
    catalog_name: str = "catalog.db",
    graph: Any = None,
    reader: Any = None,
) -> QanerisService:
    """A service wired to ``graph`` for publishing and ``reader`` for verification.

    ``graph=None`` keeps the compatibility no-op backend, which is exactly the "no verifiable
    publication" case the READY contract has to refuse. Passing a separate ``reader`` covers a store
    that publishes while the read side cannot confirm it.
    """
    return QanerisService(
        Catalog(tmp_path / catalog_name),
        graph_store=graph,
        graph_reader=reader if reader is not None else graph,
    )


def artifact_root(tmp_path: Path) -> Path:
    return tmp_path / "artifacts" / "ingestion" / "excel"


def artifact_directory_for(tmp_path: Path, data: bytes, policy: ExcelIngestionPolicy) -> Path:
    """The content-addressed directory the service derives for these bytes under this policy."""
    digest = hashlib.sha256(data).hexdigest()
    workspace = hashlib.sha256(b"default").hexdigest()[:12]
    return artifact_root(tmp_path) / workspace / artifact_key(policy) / digest


def import_workbook(
    tmp_path: Path,
    data: bytes,
    *,
    policy: ExcelIngestionPolicy | None = None,
    filename: str = "orders.xlsx",
    datasource_name: str | None = None,
    app: QanerisService | None = None,
) -> tuple[Any, Path, QanerisService]:
    source = write_source(tmp_path, data, filename)
    service = app or application(tmp_path, graph=FakeGraph())
    service_ingestion = ingestion(tmp_path, service, policy=policy)
    result = service_ingestion.import_excel(
        ExcelImportRequest(
            file_path=str(source), datasource_name=datasource_name, workspace_id="default"
        )
    )
    return result, source, service


def expect(error_code: ExcelErrorCode, callable_: Any) -> ExcelIngestionError:
    with pytest.raises(ExcelIngestionError) as captured:
        callable_()
    assert captured.value.error_code == error_code.value
    return captured.value


def manifest_of(result: Any) -> dict[str, Any]:
    return json.loads(Path(result.manifest_path).read_text("utf-8"))


# --------------------------------------------------------------------------------------------
# container preflight
# --------------------------------------------------------------------------------------------


@pytest.mark.parametrize("name", ["orders.xls", "orders.xlsm", "orders.xlsb", "orders.ods", "orders.csv", "orders"])
def test_invalid_extension_is_rejected(tmp_path, name: str) -> None:
    data = build_workbook({"orders": simple_rows()})
    expect(
        ExcelErrorCode.INVALID_EXTENSION,
        lambda: read_container(data, name),
    )


def test_corrupt_bytes_are_rejected() -> None:
    expect(
        ExcelErrorCode.INVALID_CONTAINER,
        lambda: read_container(b"this is not a zip archive", "orders.xlsx"),
    )


def test_non_spreadsheet_zip_is_rejected() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a spreadsheet")
    expect(
        ExcelErrorCode.INVALID_CONTAINER,
        lambda: read_container(buffer.getvalue(), "orders.xlsx"),
    )


def test_macro_container_is_rejected() -> None:
    original = build_workbook({"orders": simple_rows()})
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("xl/vbaProject.bin", b"macro")
    expect(
        ExcelErrorCode.INVALID_CONTAINER,
        lambda: read_container(buffer.getvalue(), "orders.xlsx"),
    )


def test_file_size_limit(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    policy = ExcelIngestionPolicy(max_file_size_bytes=64)
    expect(ExcelErrorCode.FILE_TOO_LARGE, lambda: read_container(data, "orders.xlsx", policy))


def test_zip_entry_limit() -> None:
    data = build_workbook({"orders": simple_rows()})
    policy = ExcelIngestionPolicy(max_zip_entries=3)
    expect(ExcelErrorCode.ZIP_LIMIT_EXCEEDED, lambda: read_container(data, "orders.xlsx", policy))


def test_zip_uncompressed_limit() -> None:
    data = build_workbook({"orders": simple_rows()})
    policy = ExcelIngestionPolicy(max_uncompressed_size_bytes=128)
    expect(
        ExcelErrorCode.ZIP_LIMIT_EXCEEDED, lambda: read_container(data, "orders.xlsx", policy)
    )


def test_container_reports_measured_limits() -> None:
    data = build_workbook({"orders": simple_rows()})
    info = read_container(data, "orders.xlsx")
    assert info.file_size == len(data)
    assert info.zip_entries > 0
    assert info.uncompressed_size >= info.file_size


def test_default_policy_limits_are_locked() -> None:
    policy = DEFAULT_EXCEL_POLICY
    assert policy.version == "excel-ingestion-v1"
    assert policy.max_file_size_bytes == 25 * 1024 * 1024
    assert policy.max_visible_sheets == 32
    assert policy.max_rows_per_sheet == 100_000
    assert policy.max_columns_per_sheet == 256
    assert policy.max_total_data_cells == 1_000_000
    assert policy.max_text_length == 32_768
    assert policy.max_zip_entries == 10_000
    assert policy.max_uncompressed_size_bytes == 200 * 1024 * 1024
    assert policy.max_column_name_length == 128


@pytest.mark.parametrize("field", ["max_visible_sheets", "max_rows_per_sheet", "max_text_length"])
def test_policy_rejects_non_positive_limits(field: str) -> None:
    with pytest.raises(ValueError, match=field):
        ExcelIngestionPolicy(**{field: 0})


# --------------------------------------------------------------------------------------------
# sheet rules
# --------------------------------------------------------------------------------------------


def test_hidden_and_very_hidden_sheets_are_skipped() -> None:
    data = build_workbook(
        {
            "orders": simple_rows(),
            "secret": [["a"], [1]],
            "draft": [["b"], [2]],
        },
        sheet_states={"secret": "hidden", "draft": "veryHidden"},
    )
    plan = validate_workbook(data)
    assert [sheet.table_name for sheet in plan.sheets] == ["orders"]
    assert len(plan.warnings) == 2
    assert all("跳过非可见工作表" in warning for warning in plan.warnings)


def test_empty_sheet_is_skipped() -> None:
    data = build_workbook({"empty": [], "orders": simple_rows()}, order=["empty", "orders"])
    plan = validate_workbook(data)
    assert [sheet.table_name for sheet in plan.sheets] == ["orders"]
    assert any("跳过空工作表" in warning for warning in plan.warnings)


def test_no_importable_sheet_fails() -> None:
    data = build_workbook({"empty": [], "hidden_only": [["a"], [1]]}, sheet_states={"hidden_only": "hidden"})
    expect(ExcelErrorCode.NO_IMPORTABLE_SHEET, lambda: validate_workbook(data))


def test_visible_sheet_limit() -> None:
    data = build_workbook({"a1": [["x"], [1]], "a2": [["x"], [1]]})
    policy = ExcelIngestionPolicy(max_visible_sheets=1)
    expect(ExcelErrorCode.SHEET_LIMIT_EXCEEDED, lambda: validate_workbook(data, policy))


def test_merged_range_fails_the_whole_workbook() -> None:
    data = build_workbook(
        {"orders": simple_rows(), "merged_sheet": [["a", "b"], [1, 2]]},
        merges={"merged_sheet": ["A1:B1"]},
    )
    error = expect(ExcelErrorCode.MERGED_CELL_UNSUPPORTED, lambda: validate_workbook(data))
    assert error.sheet == "merged_sheet"


def test_formula_fails_the_whole_workbook_without_echoing_the_formula() -> None:
    data = build_workbook(
        {"orders": simple_rows(), "calc": [["total"], ["=SUM(orders!A1:A2)"]]}
    )
    error = expect(ExcelErrorCode.FORMULA_UNSUPPORTED, lambda: validate_workbook(data))
    assert error.sheet == "calc"
    assert "SUM" not in str(error)
    assert error.coordinate == "A2"


def test_cached_formula_result_is_not_used_as_a_value() -> None:
    data = build_workbook({"orders": [["a"], ["=1+1"]]})
    expect(ExcelErrorCode.FORMULA_UNSUPPORTED, lambda: validate_workbook(data))


def test_row_limit() -> None:
    data = build_workbook({"orders": [["a"], [1], [2]]})
    policy = ExcelIngestionPolicy(max_rows_per_sheet=2)
    expect(ExcelErrorCode.ROW_LIMIT_EXCEEDED, lambda: validate_workbook(data, policy))


def test_column_limit() -> None:
    data = build_workbook({"orders": [["a", "b", "c"], [1, 2, 3]]})
    policy = ExcelIngestionPolicy(max_columns_per_sheet=2)
    expect(ExcelErrorCode.COLUMN_LIMIT_EXCEEDED, lambda: validate_workbook(data, policy))


def test_cell_limit() -> None:
    data = build_workbook({"orders": [["a", "b"], [1, 2], [3, 4]]})
    policy = ExcelIngestionPolicy(max_total_data_cells=3)
    expect(ExcelErrorCode.CELL_LIMIT_EXCEEDED, lambda: validate_workbook(data, policy))


def test_text_length_limit() -> None:
    data = build_workbook({"orders": [["memo"], ["x" * 9]]})
    policy = ExcelIngestionPolicy(max_text_length=8)
    error = expect(
        ExcelErrorCode.INVALID_CELL_VALUE, lambda: validate_workbook(data, policy)
    )
    assert error.coordinate == "A2"


def test_header_must_be_the_first_row() -> None:
    data = build_workbook({"orders": [[], ["title"], ["order_id"], [1]]})
    expect(ExcelErrorCode.INVALID_HEADER, lambda: validate_workbook(data))


def test_empty_header_fails() -> None:
    data = build_workbook({"orders": [["order_id", None, "region"], [1, 2, 3]]})
    error = expect(ExcelErrorCode.INVALID_HEADER, lambda: validate_workbook(data))
    assert error.coordinate == "B1"


def test_duplicate_header_is_case_insensitive() -> None:
    data = build_workbook({"orders": [["Amount", "amount"], [1, 2]]})
    error = expect(ExcelErrorCode.DUPLICATE_HEADER, lambda: validate_workbook(data))
    assert error.coordinate == "B1"


def test_header_length_limit() -> None:
    data = build_workbook({"orders": [["x" * 9], [1]]})
    policy = ExcelIngestionPolicy(max_column_name_length=8)
    expect(ExcelErrorCode.INVALID_HEADER, lambda: validate_workbook(data, policy))


def test_header_cannot_contain_nul() -> None:
    """OOXML cannot carry a NUL, so the guard is exercised at the rule boundary directly."""
    from qaneris.ingestion.excel.validator import _parse_header, _StoredCell, _table_name

    header, error = _parse_header(
        [_StoredCell(column=1, value="or\x00der_id", data_type="s", number_format=None)],
        1,
        "orders",
        DEFAULT_EXCEL_POLICY,
    )
    assert header == {}
    assert error is not None
    assert error.error_code == ExcelErrorCode.INVALID_HEADER.value

    with pytest.raises(ExcelIngestionError) as captured:
        _table_name("or\x00der", set())
    assert captured.value.error_code == ExcelErrorCode.INVALID_TABLE_NAME.value


def test_duplicate_table_name_after_normalisation() -> None:
    data = build_workbook({"①": [["a"], [1]], "1": [["b"], [2]]})
    expect(ExcelErrorCode.DUPLICATE_TABLE_NAME, lambda: validate_workbook(data))


def test_data_cell_beyond_the_header_range_fails() -> None:
    data = build_workbook({"orders": [["order_id"], [1, "stray"]]})
    error = expect(ExcelErrorCode.INVALID_HEADER, lambda: validate_workbook(data))
    assert error.coordinate == "B2"


def test_unsupported_cell_value_fails() -> None:
    data = build_workbook({"orders": [["span"], [dt.timedelta(days=1)]]})
    expect(ExcelErrorCode.INVALID_CELL_VALUE, lambda: validate_workbook(data))


def test_empty_rows_are_skipped() -> None:
    data = build_workbook({"orders": [["order_id", "region"], [None, None], [1, "east"], [None, None]]})
    plan = validate_workbook(data)
    assert plan.sheets[0].row_count == 1


def test_sheet_name_is_normalised_without_translation() -> None:
    data = build_workbook({" 订单 ① ": [["金额"], [1]]})
    plan = validate_workbook(data)
    assert plan.sheets[0].original_name == " 订单 ① "
    assert plan.sheets[0].table_name == "订单 1"


def test_column_name_is_normalised_and_keeps_its_language() -> None:
    data = build_workbook({"orders": [["  订单 ①  ", "AMOUNT"], [1, 2]]})
    plan = validate_workbook(data)
    columns = plan.sheets[0].columns
    assert columns[0].source_header == "  订单 ①  "
    assert columns[0].column_name == "订单 1"
    assert columns[1].column_name == "AMOUNT"


# --------------------------------------------------------------------------------------------
# type lattice
# --------------------------------------------------------------------------------------------


def test_classify_each_single_cell_kind() -> None:
    assert classify_cell(None) is CellKind.NULL
    assert classify_cell(True) is CellKind.BOOLEAN
    assert classify_cell(7) is CellKind.INTEGER
    assert classify_cell(1.5) is CellKind.REAL
    assert classify_cell(dt.date(2024, 1, 2)) is CellKind.DATE
    assert classify_cell(naive_datetime(2024, 1, 2, 3, 4, 5)) is CellKind.DATETIME
    assert classify_cell(dt.time(3, 4, 5)) is CellKind.TIME
    assert classify_cell("text") is CellKind.TEXT


def test_classify_splits_date_from_datetime_by_stored_format() -> None:
    midnight = naive_datetime(2024, 1, 2)
    assert classify_cell(midnight, "yyyy-mm-dd") is CellKind.DATE
    assert classify_cell(midnight, "yyyy-mm-dd hh:mm:ss") is CellKind.DATETIME
    assert classify_cell(naive_datetime(2024, 1, 2, 3, 4, 5), "yyyy-mm-dd") is CellKind.DATETIME


@pytest.mark.parametrize("value", [float("nan"), float("inf"), float("-inf")])
def test_classify_rejects_non_finite_reals(value: float) -> None:
    with pytest.raises(UnsupportedCellValue):
        classify_cell(value)


@pytest.mark.parametrize("value", [dt.timedelta(days=1), b"bytes", object(), [1, 2], {"a": 1}])
def test_classify_rejects_complex_values(value: Any) -> None:
    with pytest.raises(UnsupportedCellValue):
        classify_cell(value)


@pytest.mark.parametrize(
    ("number_format", "expected"),
    [
        (None, False),
        ("General", False),
        ("0.00", False),
        ("yyyy-mm-dd", False),
        ("yyyy\\-mm\\-dd", False),
        ("mm/dd/yyyy", False),
        ('"h"0.00', False),
        ("hh:mm:ss", True),
        ("yyyy-mm-dd hh:mm:ss", True),
        ("[h]:mm:ss", True),
        ("m/d/yy h:mm AM/PM", True),
    ],
)
def test_format_has_time(number_format: str | None, expected: bool) -> None:
    from qaneris.ingestion.excel.types import format_has_time

    assert format_has_time(number_format) is expected


@pytest.mark.parametrize(
    ("left", "right", "expected"),
    [
        (None, ColumnType.INTEGER, ColumnType.INTEGER),
        (ColumnType.INTEGER, None, ColumnType.INTEGER),
        (ColumnType.BOOLEAN, ColumnType.BOOLEAN, ColumnType.BOOLEAN),
        (ColumnType.INTEGER, ColumnType.REAL, ColumnType.REAL),
        (ColumnType.DATE, ColumnType.DATETIME, ColumnType.DATETIME),
        (ColumnType.INTEGER, ColumnType.TEXT, ColumnType.TEXT),
        (ColumnType.DATE, ColumnType.TEXT, ColumnType.TEXT),
        (ColumnType.BOOLEAN, ColumnType.INTEGER, ColumnType.TEXT),
        (ColumnType.BOOLEAN, ColumnType.REAL, ColumnType.TEXT),
        (ColumnType.DATE, ColumnType.INTEGER, ColumnType.TEXT),
        (ColumnType.DATE, ColumnType.REAL, ColumnType.TEXT),
        (ColumnType.DATE, ColumnType.BOOLEAN, ColumnType.TEXT),
        (ColumnType.DATETIME, ColumnType.REAL, ColumnType.TEXT),
    ],
)
def test_widening_table(left, right, expected) -> None:
    from qaneris.ingestion.excel.types import widen

    assert widen(left, right) == expected
    assert widen(right, left) == expected


def test_infer_column_type_maps_kinds_deterministically() -> None:
    assert infer_column_type([]) == ColumnType.TEXT
    assert infer_column_type([CellKind.NULL]) == ColumnType.TEXT
    assert infer_column_type([CellKind.BOOLEAN]) == ColumnType.BOOLEAN
    assert infer_column_type([CellKind.INTEGER]) == ColumnType.INTEGER
    assert infer_column_type([CellKind.REAL]) == ColumnType.REAL
    assert infer_column_type([CellKind.DATE]) == ColumnType.DATE
    assert infer_column_type([CellKind.DATETIME]) == ColumnType.DATETIME
    assert infer_column_type([CellKind.TIME]) == ColumnType.TEXT
    assert infer_column_type([CellKind.INTEGER, CellKind.REAL]) == ColumnType.REAL
    assert infer_column_type([CellKind.DATE, CellKind.DATETIME]) == ColumnType.DATETIME
    assert infer_column_type([CellKind.INTEGER, CellKind.TEXT]) == ColumnType.TEXT
    assert infer_column_type([CellKind.BOOLEAN, CellKind.INTEGER]) == ColumnType.TEXT
    assert infer_column_type([CellKind.DATE, CellKind.INTEGER]) == ColumnType.TEXT


def test_workbook_infers_every_declared_type() -> None:
    data = build_workbook(
        {
            "orders": [
                ["flag", "count", "ratio", "day", "moment", "label"],
                [True, 1, 1.5, dt.date(2024, 1, 2), naive_datetime(2024, 1, 2, 3, 4, 5), "east"],
            ]
        },
        number_formats={"orders": {"D2": "yyyy-mm-dd", "E2": "yyyy-mm-dd hh:mm:ss"}},
    )
    columns = {column.column_name: column for column in validate_workbook(data).sheets[0].columns}
    assert [columns[name].inferred_type for name in ("flag", "count", "ratio", "day", "moment", "label")] == [
        ColumnType.BOOLEAN,
        ColumnType.INTEGER,
        ColumnType.REAL,
        ColumnType.DATE,
        ColumnType.DATETIME,
        ColumnType.TEXT,
    ]


def test_integer_and_real_widen_to_real() -> None:
    data = build_workbook({"orders": [["amount"], [1.5], [2]]})
    column = validate_workbook(data).sheets[0].columns[0]
    assert column.inferred_type == ColumnType.REAL
    assert column.all_null is False


def test_date_and_datetime_widen_to_datetime() -> None:
    data = build_workbook(
        {"orders": [["when"], [dt.date(2024, 1, 2)], [naive_datetime(2024, 1, 3, 4, 5, 6)]]},
        number_formats={"orders": {"A2": "yyyy-mm-dd", "A3": "yyyy-mm-dd hh:mm:ss"}},
    )
    column = validate_workbook(data).sheets[0].columns[0]
    assert column.inferred_type == ColumnType.DATETIME


def test_any_type_with_text_widens_to_text() -> None:
    data = build_workbook({"orders": [["value"], [1], ["north"]]})
    column = validate_workbook(data).sheets[0].columns[0]
    assert column.inferred_type == ColumnType.TEXT


def test_boolean_with_number_widens_to_text() -> None:
    data = build_workbook({"orders": [["value"], [True], [5]]})
    assert validate_workbook(data).sheets[0].columns[0].inferred_type == ColumnType.TEXT


def test_date_with_number_widens_to_text() -> None:
    data = build_workbook(
        {"orders": [["value"], [dt.date(2024, 1, 2)], [5]]},
        number_formats={"orders": {"A2": "yyyy-mm-dd"}},
    )
    column = validate_workbook(data).sheets[0].columns[0]
    assert column.inferred_type == ColumnType.TEXT


def test_all_null_column_is_text_and_marked() -> None:
    data = build_workbook({"orders": [["used", "unused"], [1, None], [2, None]]})
    columns = {column.column_name: column for column in validate_workbook(data).sheets[0].columns}
    assert columns["unused"].inferred_type == ColumnType.TEXT
    assert columns["unused"].all_null is True
    assert columns["unused"].nullable is True
    assert columns["used"].nullable is False


def test_nullable_is_true_when_a_data_cell_is_missing() -> None:
    data = build_workbook({"orders": [["a", "b"], [1, 2], [3, None]]})
    columns = {column.column_name: column for column in validate_workbook(data).sheets[0].columns}
    assert columns["b"].nullable is True
    assert columns["b"].all_null is False


def test_coerce_cell_maps_every_declared_type() -> None:
    assert coerce_cell(None, None, ColumnType.TEXT) is None
    assert coerce_cell(True, None, ColumnType.BOOLEAN) == 1
    assert coerce_cell(False, None, ColumnType.BOOLEAN) == 0
    assert coerce_cell(5, None, ColumnType.INTEGER) == 5
    assert coerce_cell(5, None, ColumnType.REAL) == 5.0
    assert coerce_cell(dt.date(2024, 1, 2), None, ColumnType.DATE) == "2024-01-02"
    assert (
        coerce_cell(naive_datetime(2024, 1, 2, 3, 4, 5), "hh:mm", ColumnType.DATETIME)
        == "2024-01-02T03:04:05"
    )
    assert coerce_cell(naive_datetime(2024, 1, 2), "yyyy-mm-dd", ColumnType.TEXT) == "2024-01-02"
    assert coerce_cell(True, None, ColumnType.TEXT) == "true"
    assert coerce_cell(dt.time(1, 2, 3), None, ColumnType.TEXT) == "01:02:03"
    assert coerce_cell("east", None, ColumnType.TEXT) == "east"


# --------------------------------------------------------------------------------------------
# materialization
# --------------------------------------------------------------------------------------------


def _materialize(tmp_path: Path, data: bytes, destination_name: str = "data.sqlite3") -> tuple[Path, Any]:
    plan = validate_workbook(data)
    destination = tmp_path / destination_name
    tables = [
        TableWrite(
            table_name=sheet.table_name,
            columns=sheet.columns,
            row_count=sheet.row_count,
            rows=(lambda selected=sheet: iter_table_rows(plan, selected)),
        )
        for sheet in plan.sheets
    ]
    SQLiteMaterializer().materialize(tables, destination)
    return destination, plan


def test_quote_identifier_escapes_embedded_quotes() -> None:
    assert quote_identifier("orders") == '"orders"'
    assert quote_identifier('we"ird') == '"we""ird"'
    with pytest.raises(ValueError):
        quote_identifier("")
    with pytest.raises(ValueError):
        quote_identifier("a\x00b")


def test_declared_types_and_stored_values(tmp_path) -> None:
    data = build_workbook(
        {
            "orders": [
                ["flag", "count", "ratio", "day", "moment", "label"],
                [True, 1, 1.5, dt.date(2024, 1, 2), naive_datetime(2024, 1, 2, 3, 4, 5), "east"],
            ]
        },
        number_formats={"orders": {"D2": "yyyy-mm-dd", "E2": "yyyy-mm-dd hh:mm:ss"}},
    )
    destination, _ = _materialize(tmp_path, data)
    with sqlite3.connect(f"file:{destination}?mode=ro", uri=True) as connection:
        declared = {
            row[1]: row[2] for row in connection.execute('PRAGMA table_info("orders")')
        }
        row = connection.execute('SELECT * FROM "orders"').fetchone()
    assert declared == {
        "flag": "BOOLEAN",
        "count": "INTEGER",
        "ratio": "REAL",
        "day": "DATE",
        "moment": "DATETIME",
        "label": "TEXT",
    }
    assert row == (1, 1, 1.5, "2024-01-02", "2024-01-02T03:04:05", "east")


def test_identifiers_with_quotes_round_trip(tmp_path) -> None:
    data = build_workbook({'a"b': [['c"d'], ["v"]]})
    destination, plan = _materialize(tmp_path, data)
    assert plan.sheets[0].table_name == 'a"b'
    with sqlite3.connect(f"file:{destination}?mode=ro", uri=True) as connection:
        assert connection.execute('SELECT "c""d" FROM "a""b"').fetchall() == [("v",)]


def test_values_are_bound_as_parameters_not_concatenated(tmp_path) -> None:
    hostile = "x'); DROP TABLE \"orders\"; --"
    data = build_workbook({"orders": [["memo"], [hostile]]})
    destination, _ = _materialize(tmp_path, data)
    with sqlite3.connect(f"file:{destination}?mode=ro", uri=True) as connection:
        assert connection.execute('SELECT memo FROM "orders"').fetchall() == [(hostile,)]
        assert connection.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='orders'"
        ).fetchone() == (1,)


def test_staging_is_cleaned_up_when_materialization_fails(tmp_path) -> None:
    columns = (
        ColumnPlan(
            source_header="a",
            column_name="a",
            inferred_type=ColumnType.INTEGER,
            nullable=False,
            all_null=False,
        ),
    )
    table = TableWrite(
        table_name="orders", columns=columns, row_count=5, rows=lambda: iter([(1,), (2,)])
    )
    destination = tmp_path / "data.sqlite3"
    expect(
        ExcelErrorCode.MATERIALIZATION_FAILED,
        lambda: SQLiteMaterializer().materialize([table], destination),
    )
    assert not destination.exists()
    assert sorted(path.name for path in tmp_path.iterdir()) == []


def test_staging_is_cleaned_up_when_a_row_is_rejected(tmp_path) -> None:
    columns = (
        ColumnPlan(
            source_header="a",
            column_name="a",
            inferred_type=ColumnType.INTEGER,
            nullable=False,
            all_null=False,
        ),
    )
    table = TableWrite(
        table_name="orders", columns=columns, row_count=1, rows=lambda: iter([(1, 2)])
    )
    expect(
        ExcelErrorCode.MATERIALIZATION_FAILED,
        lambda: SQLiteMaterializer().materialize([table], tmp_path / "data.sqlite3"),
    )
    assert sorted(path.name for path in tmp_path.iterdir()) == []


def test_materialized_file_is_private_and_has_no_journal_leftovers(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    destination, _ = _materialize(tmp_path, data)
    assert destination.stat().st_mode & 0o777 == 0o600
    assert sorted(path.name for path in tmp_path.iterdir()) == ["data.sqlite3"]


def test_is_finalized_is_an_observation_not_an_assumption(tmp_path) -> None:
    """The reuse check only accepts a complete, readable database of the expected shape."""
    materializer = SQLiteMaterializer()
    missing = tmp_path / "absent.sqlite3"
    assert materializer.is_finalized(missing, ["orders"]) is False

    destination, _ = _materialize(tmp_path, build_workbook({"orders": simple_rows()}))
    assert materializer.is_finalized(destination, ["orders"]) is True
    assert materializer.is_finalized(destination, ["orders", "customers"]) is False
    assert materializer.is_finalized(destination, ["customers"]) is False

    unreadable = tmp_path / "truncated.sqlite3"
    unreadable.write_bytes(destination.read_bytes()[:64])
    assert materializer.is_finalized(unreadable, ["orders"]) is False


# --------------------------------------------------------------------------------------------
# service: artifact, idempotency, failure record
# --------------------------------------------------------------------------------------------


def test_import_reaches_ready_and_writes_the_managed_layout(tmp_path) -> None:
    graph = FakeGraph()
    result, source, service = import_workbook(
        tmp_path, build_workbook({"orders": simple_rows()}), app=application(tmp_path, graph=graph)
    )
    directory = Path(result.artifact_directory)
    assert result.status == "READY"
    assert sorted(path.name for path in directory.iterdir()) == [
        "data.sqlite3",
        "manifest.json",
        "source.xlsx",
    ]
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    assert result.file_sha256 == digest
    # <root>/ingestion/excel/<workspace_digest>/<policy key>/<file sha256>
    assert directory.name == digest
    assert directory.parent.name == artifact_key(DEFAULT_EXCEL_POLICY)
    assert directory.parent.parent.name == hashlib.sha256(b"default").hexdigest()[:12]
    assert result.import_id == (
        f"excel_{directory.parent.parent.name}_{policy_digest(DEFAULT_EXCEL_POLICY)}_{digest[:16]}"
    )
    assert Path(result.sqlite_path) == directory / "data.sqlite3"
    assert result.workspace_id == "default"
    assert result.datasource_id.startswith("ds_")
    assert result.snapshot_id.startswith("snap_")
    assert result.scan_version == 1
    # A READY import is only returned once the publication was really observed.
    assert result.neo4j_publication_verified is True
    assert graph.published[("default", result.datasource_id)] == result.scan_version
    datasource, _ = service.catalog.get_datasource(result.datasource_id)
    assert datasource.status == "ready"
    assert datasource.driver == "sqlite"
    profile = service.catalog.get_connection_profile(result.datasource_id)
    assert profile is not None
    assert profile.endpoint.path == result.sqlite_path
    assert [sheet.table_name for sheet in result.sheets] == ["orders"]
    assert result.sheets[0].row_count == 2


def test_import_scans_through_an_explicit_initialize_not_the_create(
    tmp_path, monkeypatch
) -> None:
    """``Test → Save → Scan`` in the secure contract means an import must call the scan step itself.

    ``create_secure_datasource`` saves without scanning, so an import that relied on create to reach
    READY would either stop working or, worse, relax the READY definition. This pins the explicit
    ``initialize_datasource`` call: if a future change removes it, the import cannot reach READY
    through this test.
    """
    graph = FakeGraph()
    service = application(tmp_path, graph=graph)
    calls: list[str] = []
    original = service.initialize_datasource

    def tracked(datasource_id: str):
        calls.append(datasource_id)
        return original(datasource_id)

    monkeypatch.setattr(service, "initialize_datasource", tracked)

    result, _, _ = import_workbook(
        tmp_path, build_workbook({"orders": simple_rows()}), app=service
    )

    assert calls == [result.datasource_id]
    assert result.status == "READY"


def test_manifest_records_shape_but_never_business_cell_values(tmp_path) -> None:
    secret = "SUPER-SECRET-CUSTOMER-42"
    data = build_workbook({"orders": [["order_id", "customer"], [1, secret]]})
    result, _, _ = import_workbook(tmp_path, data)
    manifest_text = Path(result.manifest_path).read_text("utf-8")
    manifest = json.loads(manifest_text)
    assert secret not in manifest_text
    assert manifest["format_version"] == "excel-manifest-v1"
    assert manifest["policy_version"] == "excel-ingestion-v1"
    assert manifest["status"] == "READY"
    assert manifest["error_code"] is None
    assert manifest["sheets"][0]["columns"][1] == {
        "source_header": "customer",
        "column_name": "customer",
        "inferred_type": "TEXT",
        "nullable": False,
        "all_null": False,
    }
    assert manifest["datasource_id"] == result.datasource_id
    assert manifest["snapshot_id"] == result.snapshot_id
    assert manifest["scan_version"] == result.scan_version
    # The value really was imported, it is only absent from the lifecycle record.
    with sqlite3.connect(f"file:{result.sqlite_path}?mode=ro", uri=True) as connection:
        assert connection.execute('SELECT customer FROM "orders"').fetchone() == (secret,)


def test_same_bytes_are_idempotent(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    service = application(tmp_path, graph=FakeGraph())
    state = ingestion(tmp_path, service)
    request = ExcelImportRequest(file_path=str(source))

    first = state.import_excel(request)
    sqlite_path = Path(first.sqlite_path)
    before = (sqlite_path.stat().st_ino, sqlite_path.stat().st_mtime_ns, sqlite_path.stat().st_size)

    second = state.import_excel(request)

    assert second.import_id == first.import_id
    assert second.datasource_id == first.datasource_id
    assert second.snapshot_id == first.snapshot_id
    assert second.scan_version == first.scan_version
    assert second.neo4j_publication_verified is True
    assert (sqlite_path.stat().st_ino, sqlite_path.stat().st_mtime_ns, sqlite_path.stat().st_size) == before
    assert len(service.catalog.list_datasources("default")) == 1
    with sqlite3.connect(service.catalog.path) as connection:
        snapshots = connection.execute(
            "SELECT COUNT(*) FROM scan_snapshot WHERE datasource_id=?", (first.datasource_id,)
        ).fetchone()[0]
    assert snapshots == 1


def test_different_bytes_are_a_new_import(tmp_path) -> None:
    service = application(tmp_path, graph=FakeGraph())
    first, _, _ = import_workbook(
        tmp_path, build_workbook({"orders": simple_rows()}), app=service, filename="a.xlsx"
    )
    second, _, _ = import_workbook(
        tmp_path,
        build_workbook({"orders": [["order_id", "region"], [9, "north"]]}),
        app=service,
        filename="b.xlsx",
    )
    assert second.import_id != first.import_id
    assert second.datasource_id != first.datasource_id
    assert Path(second.artifact_directory) != Path(first.artifact_directory)
    assert len(service.catalog.list_datasources("default")) == 2


def test_ready_import_does_not_rewrite_the_source_workbook(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    service = application(tmp_path, graph=FakeGraph())
    state = ingestion(tmp_path, service)
    request = ExcelImportRequest(file_path=str(source))

    first = state.import_excel(request)
    managed = Path(first.artifact_directory) / "source.xlsx"
    before = (managed.stat().st_ino, managed.stat().st_mtime_ns)
    state.import_excel(request)

    assert (managed.stat().st_ino, managed.stat().st_mtime_ns) == before
    assert managed.read_bytes() == data


def test_failed_validation_records_a_failed_manifest_and_keeps_the_source(tmp_path) -> None:
    data = build_workbook({"orders": [["total"], ["=1+1"]]})
    source = write_source(tmp_path, data)
    service = application(tmp_path)
    state = ingestion(tmp_path, service)

    expect(
        ExcelErrorCode.FORMULA_UNSUPPORTED,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(source))),
    )

    directory = artifact_directory_for(tmp_path, data, DEFAULT_EXCEL_POLICY)
    manifest = json.loads((directory / "manifest.json").read_text("utf-8"))
    assert manifest["status"] == "FAILED"
    assert manifest["error_code"] == "EXCEL_FORMULA_UNSUPPORTED"
    assert manifest["datasource_id"] is None
    assert manifest["snapshot_id"] is None
    assert (directory / "source.xlsx").read_bytes() == data
    assert not (directory / "data.sqlite3").exists()
    assert service.catalog.list_datasources("default") == []


def test_container_failure_never_creates_the_artifact_directory(tmp_path) -> None:
    source = write_source(tmp_path, b"not a workbook", "orders.csv")
    service = application(tmp_path)
    state = ingestion(tmp_path, service)
    expect(
        ExcelErrorCode.INVALID_EXTENSION,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(source))),
    )
    assert not (tmp_path / "artifacts").exists()


def test_missing_file_is_rejected_with_a_stable_code(tmp_path) -> None:
    state = ingestion(tmp_path, application(tmp_path))
    expect(
        ExcelErrorCode.INVALID_CONTAINER,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(tmp_path / "absent.xlsx"))),
    )


def test_import_directory_and_artifacts_are_private(tmp_path) -> None:
    result, _, _ = import_workbook(tmp_path, build_workbook({"orders": simple_rows()}))
    directory = Path(result.artifact_directory)
    assert directory.stat().st_mode & 0o777 == 0o700
    assert directory.parent.stat().st_mode & 0o777 == 0o700
    assert (directory / "source.xlsx").stat().st_mode & 0o777 == 0o600
    assert (directory / "data.sqlite3").stat().st_mode & 0o777 == 0o600
    assert (directory / "manifest.json").stat().st_mode & 0o777 == 0o600


def test_datasource_name_defaults_to_the_original_filename(tmp_path) -> None:
    result, _, service = import_workbook(tmp_path, build_workbook({"orders": simple_rows()}))
    datasource, _ = service.catalog.get_datasource(result.datasource_id)
    assert datasource.name == "Excel: orders.xlsx"


def test_explicit_datasource_name_is_used(tmp_path) -> None:
    result, _, service = import_workbook(
        tmp_path, build_workbook({"orders": simple_rows()}), datasource_name="路演订单"
    )
    datasource, _ = service.catalog.get_datasource(result.datasource_id)
    assert datasource.name == "路演订单"


def test_manifest_cache_survives_a_rewritten_manifest(tmp_path) -> None:
    """A manifest that cannot be parsed is treated as "no usable import", not as a crash."""
    result, _, _ = import_workbook(tmp_path, build_workbook({"orders": simple_rows()}))
    Path(result.manifest_path).write_text("{ not json", "utf-8")
    assert ingestion(tmp_path, application(tmp_path))._read_manifest(Path(result.manifest_path)) is None


def test_workspace_is_part_of_the_artifact_path(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    service = application(tmp_path, graph=FakeGraph())
    state = ingestion(tmp_path, service)
    result = state.import_excel(
        ExcelImportRequest(file_path=str(source), workspace_id="workspace-a")
    )
    assert result.workspace_id == "workspace-a"
    digest = hashlib.sha256(b"workspace-a").hexdigest()[:12]
    directory = Path(result.artifact_directory)
    assert directory.parent.parent.name == digest
    assert directory.parent.name == artifact_key(DEFAULT_EXCEL_POLICY)


def test_the_original_filename_is_never_a_directory_name(tmp_path) -> None:
    """Directory names are content-addressed; a business filename must not leak into the path."""
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data, name="2026 客户订单-华东.xlsx")
    service = application(tmp_path, graph=FakeGraph())
    result = ingestion(tmp_path, service).import_excel(
        ExcelImportRequest(file_path=str(source))
    )
    parts = Path(result.artifact_directory).relative_to(tmp_path / "artifacts").parts
    assert all("订单" not in part and ".xlsx" not in part for part in parts)
    assert result.original_filename == "2026 客户订单-华东.xlsx"


def test_sqlite_adapter_reads_the_materialized_file_read_only(tmp_path) -> None:
    """The finalized artifact is consumed only through the existing SQLite adapter."""
    from qaneris.adapters.relational.sqlite import SQLiteAdapter

    result, _, _ = import_workbook(tmp_path, build_workbook({"orders": simple_rows()}))
    adapter = SQLiteAdapter("ds_excel", {"driver": "sqlite", "path": result.sqlite_path})
    datasets = adapter.scan_metadata()
    assert [dataset.name for dataset in datasets] == ["orders"]
    assert adapter.scan_relations() == []
    outcome = adapter.execute('SELECT "region" FROM "orders" ORDER BY "order_id"', 10)
    assert outcome.rows == [{"region": "east"}, {"region": "west"}]
    with pytest.raises(QuerySafetyError):
        adapter.execute('DROP TABLE "orders"', 10)


# --------------------------------------------------------------------------------------------
# READY requires a verified Neo4j publication
# --------------------------------------------------------------------------------------------


def _failed_manifest(tmp_path: Path, data: bytes, policy: ExcelIngestionPolicy | None = None) -> dict:
    directory = artifact_directory_for(tmp_path, data, policy or DEFAULT_EXCEL_POLICY)
    return json.loads((directory / "manifest.json").read_text("utf-8"))


def test_without_a_graph_backend_the_import_is_never_ready(tmp_path) -> None:
    """``status=READY`` with ``neo4j_publication_verified=false`` must be impossible."""
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    service = application(tmp_path)  # compatibility no-op store and reader
    state = ingestion(tmp_path, service)

    expect(
        ExcelErrorCode.SCAN_FAILED,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(source))),
    )

    manifest = _failed_manifest(tmp_path, data)
    assert manifest["status"] == "FAILED"
    assert manifest["error_code"] == "EXCEL_SCAN_FAILED"
    # The artifact and the datasource really exist; only the publication is unconfirmed.
    assert manifest["datasource_id"] is not None
    assert manifest["snapshot_id"] is not None
    assert service.catalog.list_datasources("default") != []


def test_a_publishing_store_is_not_enough_when_the_reader_cannot_confirm(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    graph = FakeGraph()
    service = application(tmp_path, graph=graph, reader=NullGraphReader())
    state = ingestion(tmp_path, service)

    expect(
        ExcelErrorCode.SCAN_FAILED,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(source))),
    )
    assert graph.published != {}, "the store did publish; only the read side is unusable"
    assert _failed_manifest(tmp_path, data)["error_code"] == "EXCEL_SCAN_FAILED"


def test_a_publication_that_silently_drops_the_write_fails_the_import(tmp_path) -> None:
    """The scan and the snapshot look complete, so only the publication check can catch this."""
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    graph = FakeGraph()
    graph.drop_writes = True
    service = application(tmp_path, graph=graph)
    state = ingestion(tmp_path, service)

    expect(
        ExcelErrorCode.SCAN_FAILED,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(source))),
    )

    manifest = _failed_manifest(tmp_path, data)
    assert manifest["status"] == "FAILED"
    assert manifest["error_code"] == "EXCEL_SCAN_FAILED"
    assert manifest["datasource_id"] is not None  # really created, not fabricated
    snapshot = service.catalog.get_active_snapshot(manifest["datasource_id"])
    assert snapshot is not None and snapshot.status.value == "ready"


def test_an_unreachable_graph_fails_the_import(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    graph = FakeGraph()
    graph.unavailable = True
    state = ingestion(tmp_path, application(tmp_path, graph=graph))

    error = expect(
        ExcelErrorCode.SCAN_FAILED,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(source))),
    )
    assert error.error_code == ExcelErrorCode.SCAN_FAILED.value
    assert _failed_manifest(tmp_path, data)["status"] == "FAILED"


def test_a_publication_failure_inside_initialization_is_recorded(tmp_path) -> None:
    """A store that raises leaves the initialization job failed, not a READY import."""
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    service = application(tmp_path, graph=FailingGraphStore(), reader=FakeGraph())
    state = ingestion(tmp_path, service)

    expect(
        ExcelErrorCode.SCAN_FAILED,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(source))),
    )

    manifest = _failed_manifest(tmp_path, data)
    assert manifest["status"] == "FAILED"
    assert manifest["error_code"] == "EXCEL_SCAN_FAILED"
    assert manifest["snapshot_id"] is None


def test_datasource_creation_failure_records_a_failed_manifest(tmp_path, monkeypatch) -> None:
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    service = application(tmp_path, graph=FakeGraph())
    state = ingestion(tmp_path, service)

    def explode(request: Any) -> Any:
        raise RuntimeError("simulated catalog failure")

    monkeypatch.setattr(service, "create_secure_datasource", explode)

    expect(
        ExcelErrorCode.DATASOURCE_CREATION_FAILED,
        lambda: state.import_excel(ExcelImportRequest(file_path=str(source))),
    )

    manifest = _failed_manifest(tmp_path, data)
    assert manifest["status"] == "FAILED"
    assert manifest["error_code"] == "EXCEL_DATASOURCE_CREATION_FAILED"
    assert manifest["datasource_id"] is None
    assert manifest["snapshot_id"] is None


def _assert_retry_recovers(tmp_path: Path, data: bytes) -> None:
    """A failed publication must stay retryable without touching the finalized artifact."""
    source = write_source(tmp_path, data)
    graph = FakeGraph()
    graph.drop_writes = True
    service = application(tmp_path, graph=graph)
    state = ingestion(tmp_path, service)
    request = ExcelImportRequest(file_path=str(source))

    expect(ExcelErrorCode.SCAN_FAILED, lambda: state.import_excel(request))

    manifest = _failed_manifest(tmp_path, data)
    failed_datasource = manifest["datasource_id"]
    directory = artifact_directory_for(tmp_path, data, DEFAULT_EXCEL_POLICY)
    before = {
        name: (path.stat().st_ino, path.stat().st_mtime_ns, path.stat().st_size)
        for name in ("source.xlsx", "data.sqlite3")
        for path in [directory / name]
    }

    graph.drop_writes = False
    result = state.import_excel(request)

    assert result.status == "READY"
    assert result.neo4j_publication_verified is True
    # The existing datasource is reused and re-published, not duplicated.
    assert result.datasource_id == failed_datasource
    assert len(service.catalog.list_datasources("default")) == 1
    assert graph.published[("default", result.datasource_id)] == result.scan_version
    after = {
        name: (path.stat().st_ino, path.stat().st_mtime_ns, path.stat().st_size)
        for name in ("source.xlsx", "data.sqlite3")
        for path in [directory / name]
    }
    assert after == before, "a retry must not rewrite a finalized artifact"
    assert json.loads((directory / "manifest.json").read_text("utf-8"))["status"] == "READY"


def test_retry_after_a_failed_publication_recovers(tmp_path) -> None:
    _assert_retry_recovers(tmp_path, build_workbook({"orders": simple_rows()}))


def test_ready_import_is_not_reported_when_the_graph_goes_away(tmp_path) -> None:
    """A stored READY record plus an unverifiable graph must not produce a READY result."""
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    graph = FakeGraph()
    service = application(tmp_path, graph=graph)
    state = ingestion(tmp_path, service)
    request = ExcelImportRequest(file_path=str(source))

    first = state.import_excel(request)
    assert first.status == "READY"

    graph.unavailable = True
    expect(ExcelErrorCode.SCAN_FAILED, lambda: state.import_excel(request))

    # The finalized artifact keeps its own truthful record: nothing is overwritten.
    manifest = _failed_manifest(tmp_path, data)
    assert manifest["status"] == "READY"
    assert manifest["datasource_id"] == first.datasource_id
    assert manifest["snapshot_id"] == first.snapshot_id


# --------------------------------------------------------------------------------------------
# artifact immutability across policies
# --------------------------------------------------------------------------------------------


def test_a_different_policy_gets_its_own_artifact(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    service = application(tmp_path, graph=FakeGraph())
    base = ingestion(tmp_path, service)
    other_policy = replace(DEFAULT_EXCEL_POLICY, version="excel-ingestion-v2")
    other = ingestion(tmp_path, service, policy=other_policy)
    request = ExcelImportRequest(file_path=str(source), datasource_name="orders-v1")
    other_request = ExcelImportRequest(file_path=str(source), datasource_name="orders-v2")

    first = base.import_excel(request)
    first_directory = Path(first.artifact_directory)
    before = {
        name: (path.stat().st_ino, path.stat().st_mtime_ns, path.stat().st_size)
        for name in ("source.xlsx", "data.sqlite3", "manifest.json")
        for path in [first_directory / name]
    }

    second = other.import_excel(other_request)

    assert second.artifact_directory != first.artifact_directory
    assert second.import_id != first.import_id
    assert Path(second.artifact_directory).is_relative_to(tmp_path / "artifacts")
    assert Path(second.artifact_directory).parent.name == artifact_key(other_policy)
    assert Path(second.artifact_directory).parent.name != artifact_key(DEFAULT_EXCEL_POLICY)
    # A different policy is a different datasource over a different managed file.
    assert second.datasource_id != first.datasource_id
    assert len(service.catalog.list_datasources("default")) == 2
    # The old artifact is untouched, down to the inode, and keeps its READY record.
    after = {
        name: (path.stat().st_ino, path.stat().st_mtime_ns, path.stat().st_size)
        for name in ("source.xlsx", "data.sqlite3", "manifest.json")
        for path in [first_directory / name]
    }
    assert after == before
    assert json.loads((first_directory / "manifest.json").read_text("utf-8"))["status"] == "READY"
    # And the original policy is still fully idempotent.
    assert base.import_excel(request).import_id == first.import_id


def test_a_changed_limit_changes_the_artifact_identity_without_a_version_bump(tmp_path) -> None:
    data = build_workbook({"orders": simple_rows()})
    source = write_source(tmp_path, data)
    service = application(tmp_path, graph=FakeGraph())
    base = ingestion(tmp_path, service)
    tightened = replace(DEFAULT_EXCEL_POLICY, max_rows_per_sheet=10)
    assert tightened.version == DEFAULT_EXCEL_POLICY.version
    request = ExcelImportRequest(file_path=str(source), datasource_name="orders-default")
    other_request = ExcelImportRequest(file_path=str(source), datasource_name="orders-tightened")

    first = base.import_excel(request)
    second = ingestion(tmp_path, service, policy=tightened).import_excel(other_request)

    assert second.artifact_directory != first.artifact_directory
    assert policy_digest(tightened) != policy_digest(DEFAULT_EXCEL_POLICY)
    assert json.loads(
        (Path(first.artifact_directory) / "manifest.json").read_text("utf-8")
    )["status"] == "READY"


# --------------------------------------------------------------------------------------------
# application facade
# --------------------------------------------------------------------------------------------


def test_service_facade_is_the_only_entry_point(tmp_path, monkeypatch) -> None:
    """``QanerisService.import_excel`` must work end to end, not just the service class.

    The ingestion service is constructed lazily, so the facade has to hold and reuse exactly one
    instance; a missing attribute here would only surface through this entry point.
    """
    monkeypatch.setenv("QANERIS_ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    source = write_source(tmp_path, build_workbook({"orders": simple_rows()}))
    service = application(tmp_path, graph=FakeGraph())

    assert service.excel_ingestion is service.excel_ingestion

    result = service.import_excel(
        ExcelImportRequest(file_path=str(source), datasource_name="路演订单")
    )
    assert result.status == "READY"
    assert result.neo4j_publication_verified is True
    assert result.original_filename == "orders.xlsx"
    assert Path(result.artifact_directory).is_relative_to(tmp_path / "artifacts")
    assert Path(result.sqlite_path).exists()

    replay = service.import_excel(
        ExcelImportRequest(file_path=str(source), datasource_name="路演订单")
    )
    assert replay.import_id == result.import_id
    assert replay.snapshot_id == result.snapshot_id
    assert replay.datasource_id == result.datasource_id
